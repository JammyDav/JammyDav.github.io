"""
word_report_engine.py - neutral stand-in for the corporate Word engine build_docs.py expects

build_docs.py renders markdown masters through a branded engine that is not part of
this repo (and whose template carries corporate marks). This module offers the same
small surface it calls (new_report_from_template, add_heading, add_paragraph,
_add_styled_para, add_table, save, document) on a plain python-docx document in the
site's own type and colour, so the real build_docs.py runs unchanged.
"""
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

TEAL = RGBColor(0x0B, 0x6E, 0x69)
INK = RGBColor(0x1B, 0x24, 0x28)
FONT = "Segoe UI"


def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


class Report:
    def __init__(self):
        self.document = Document()
        self._document = self.document
        self._footer_metadata = {}       # accepted and ignored: no corporate footer here
        st = self.document.styles
        normal = st["Normal"]
        normal.font.name = FONT
        normal.font.size = Pt(10)
        normal.font.color.rgb = INK
        rpr = normal.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.append(fonts)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(a), FONT)
        for i, size in ((1, 20), (2, 13), (3, 11), (4, 10)):
            h = st[f"Heading {i}"]
            h.font.name = FONT
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.color.rgb = TEAL if i < 3 else INK
            h.element.get_or_add_rPr()
        for name in ("Body Text", "List Bullet"):
            st[name].font.name = FONT
        last = st.add_style("List Bullet Last", WD_STYLE_TYPE.PARAGRAPH)
        last.base_style = st["List Bullet"]
        lead = st.add_style("Lead", WD_STYLE_TYPE.PARAGRAPH)
        lead.base_style = st["Normal"]
        lead.font.size = Pt(12)
        lead.font.color.rgb = TEAL
        self._by_id = {s.style_id: s for s in st if s.style_id}
        self._by_id["lead"] = lead

    def add_heading(self, text, level=1):
        return self.document.add_heading(text, level=level)

    def add_paragraph(self, text, style=None):
        s = self._by_id.get(style) if style else None
        return self.document.add_paragraph(text, style=s)

    def _add_styled_para(self, style_id, text):
        p = self.document.add_paragraph(text, style=self._by_id.get(style_id))
        return p._p

    def add_table(self, data, headers=None, style=None, column_widths=None, auto_fit=None):
        rows = ([headers] if headers else []) + list(data)
        t = self.document.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = self.document.styles["Table Grid"]
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                cell = t.cell(r, c)
                cell.text = str(val)
                if column_widths:
                    cell.width = Inches(column_widths[c])
                if r == 0 and headers:
                    _shade(cell, "0B6E69")
                    for run in cell.paragraphs[0].runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        return t

    def save(self, path):
        self.document.save(path)


class WordReportEngine:
    def new_report_from_template(self, template=None, title=None):
        return Report()
